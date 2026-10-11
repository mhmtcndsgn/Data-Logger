"""Oturumu TXT / CSV / Excel (.xlsx) olarak dışa aktarır."""
from __future__ import annotations

import csv
import datetime as _dt
import os

from . import __version__, APP_NAME
from .i18n import T
from .parser import fmt_num
from .session import Session

FORMATS = {"txt": ".txt", "csv": ".csv", "xlsx": ".xlsx"}
XLSX_MAX_ROWS = 1_000_000          # Excel sınırı 1.048.576; güvenli pay
XLSX_CHART_MAX_ROWS = 20_000


def _columns(s: Session, lang: str, n: int):
    """(başlıklar, satır üreten fonksiyon) döndürür."""
    t = T(lang)
    heads = [t("col_index")]
    getters = [lambda i: i + 1]
    if s.has_seq and len(s.seq) >= n:
        heads.append(t("col_seq"))
        getters.append(lambda i: s.seq[i])
    if s.sample_rate > 0:
        xs, _ = s.x_values(end=n, prefer="time")
        heads.append(t("col_time_s"))
        getters.append(lambda i, xs=xs: round(xs[i], 9))
    if s.mode == "live" and len(s.t_host) >= n and n:
        base = s.started_at
        th = s.t_host
        t0 = th[0]
        heads.append(t("col_pc_time"))
        getters.append(lambda i: (base + _dt.timedelta(seconds=th[i] - t0)))
        heads.append(t("col_elapsed"))
        getters.append(lambda i: round(th[i] - t0, 4))
    for name, col in zip(s.channels, s.cols):
        heads.append(name)
        getters.append(lambda i, col=col: col[i])
    return heads, getters


def _fmt_cell(v) -> str:
    if isinstance(v, float):
        return fmt_num(v)
    if isinstance(v, _dt.datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
    return str(v)


def summary_rows(s: Session, lang: str, n: int | None = None) -> list[tuple[str, str]]:
    t = T(lang)
    n = s.n if n is None else n
    rows = [
        (t("sum_app"), f"{APP_NAME} v{__version__}"),
        (t("sum_mode"), t("mode_" + s.mode)),
    ]
    if s.port:
        rows.append((t("sum_port"), f"{s.port} @ {s.baud} baud"))
    rows += [
        (t("sum_start"), s.started_at.strftime("%Y-%m-%d %H:%M:%S")),
        (t("sum_end"), (s.ended_at or _dt.datetime.now()).strftime("%Y-%m-%d %H:%M:%S")),
        (t("sum_rows"), str(n)),
        (t("sum_channels"), ", ".join(s.channels)),
    ]
    if s.sample_rate > 0:
        rows.append((t("sum_fs"), f"{fmt_num(s.sample_rate)} Hz"))
    if s.expected is not None:
        rows.append((t("sum_expected"), str(s.expected)))
    if s.integrity_known:
        rows.append((t("sum_missing"), str(s.missing)))
        rows.append((t("sum_integrity"), t("ok") if s.integrity_ok else t("not_ok")))
    else:
        rows.append((t("sum_integrity"), t("integrity_unknown")))
    if s.gaps:
        rows.append((t("sum_gaps"), "; ".join(f"{g.after_seq}→{g.next_seq} (-{g.missing})" for g in s.gaps[:50])
                     + (" …" if len(s.gaps) > 50 else "")))
    if s.bad_lines:
        rows.append((t("sum_bad_lines"), str(s.bad_lines)))
    if s.fixed_seq:
        rows.append((t("sum_fixed"), str(s.fixed_seq)))
    if s.out_of_order:
        rows.append((t("sum_ooo"), str(s.out_of_order)))
    if s.text_lines:
        rows.append((t("sum_text_lines"), str(s.text_lines)))
    return rows


def default_filename(s: Session, prefix: str, fmt: str) -> str:
    stamp = s.started_at.strftime("%Y-%m-%d_%H-%M-%S")
    prefix = (prefix or "deney").strip().replace(" ", "_")
    return f"{prefix}_{stamp}{FORMATS[fmt]}"


def export_session(s: Session, path: str, fmt: str, lang: str = "tr") -> str:
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    with s.lock:
        if s.ended_at is not None:
            s.flush_pending()
        n = s.n          # kayıt sürerken dışa aktarılırsa: o anki satırlar
    if fmt == "xlsx":
        _export_xlsx(s, path, lang, n)
    elif fmt == "csv":
        _export_delim(s, path, lang, ",", False, n)
    else:
        _export_delim(s, path, lang, "\t", True, n)
    return path


def _export_delim(s: Session, path: str, lang: str, delim: str, comments: bool, n: int) -> None:
    heads, getters = _columns(s, lang, n)
    with open(path, "w", encoding="utf-8-sig" if not comments else "utf-8", newline="") as f:
        if comments:
            for k, v in summary_rows(s, lang, n):
                f.write(f"# {k}: {v}\n")
        w = csv.writer(f, delimiter=delim, lineterminator="\n")
        w.writerow(heads)
        for i in range(n):
            w.writerow([_fmt_cell(g(i)) for g in getters])


def _num(v):
    if isinstance(v, float):
        if v != v:
            return None
        if v.is_integer() and abs(v) < 1e15:
            return int(v)
    return v


def _export_xlsx(s: Session, path: str, lang: str, n: int) -> None:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.chart import Reference, ScatterChart, Series
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    t = T(lang)
    heads, getters = _columns(s, lang, n)
    wb = Workbook(write_only=True)

    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1F6FEB")
    center = Alignment(horizontal="center", vertical="center")

    # --- Özet sayfası ----------------------------------------------------
    ws_sum = wb.create_sheet(t("sheet_summary"))
    ws_sum.column_dimensions["A"].width = 26
    ws_sum.column_dimensions["B"].width = 36
    for col in "CDEFG":
        ws_sum.column_dimensions[col].width = 14

    def hcell(ws, v):
        c = WriteOnlyCell(ws, value=v)
        c.font = hdr_font
        c.fill = hdr_fill
        c.alignment = center
        return c

    title = WriteOnlyCell(ws_sum, value=f"{APP_NAME} — {t('sheet_summary')}")
    title.font = Font(bold=True, size=14, color="1F6FEB")
    ws_sum.append([title])
    ws_sum.append([])
    for k, v in summary_rows(s, lang, n):
        kc = WriteOnlyCell(ws_sum, value=k)
        kc.font = Font(bold=True)
        vc = WriteOnlyCell(ws_sum, value=int(v) if v.isdigit() else v)
        if k == t("sum_integrity"):
            ok = v == t("ok")
            vc.font = Font(bold=True, color="1A7F37" if ok else ("CF222E" if v == t("not_ok") else "9A6700"))
        ws_sum.append([kc, vc])
    ws_sum.append([])
    ws_sum.append([hcell(ws_sum, h) for h in (t("stat_channel"), t("stat_count"), "Min", "Max",
                                              t("stat_mean"), t("stat_std"))])
    for st in s.channel_stats(n):
        if st["count"]:
            ws_sum.append([st["name"], st["count"], _num(st["min"]), _num(st["max"]),
                           round(st["mean"], 6), round(st["std"], 6)])
        else:
            ws_sum.append([st["name"], 0])

    # --- Veri sayfaları ----------------------------------------------------
    data_sheets = []
    n_parts = max(1, -(-n // XLSX_MAX_ROWS))
    for part in range(n_parts):
        name = t("sheet_data") if n_parts == 1 else f"{t('sheet_data')} {part + 1}"
        ws = wb.create_sheet(name)
        ws.freeze_panes = "A2"
        for ci, h in enumerate(heads, start=1):
            ws.column_dimensions[get_column_letter(ci)].width = max(12, min(28, len(str(h)) + 4))
        ws.append([hcell(ws, h) for h in heads])
        lo, hi = part * XLSX_MAX_ROWS, min(n, (part + 1) * XLSX_MAX_ROWS)
        for i in range(lo, hi):
            ws.append([_num(g(i)) for g in getters])
        data_sheets.append((ws, hi - lo))

    # --- Grafik ------------------------------------------------------------
    if 0 < n <= XLSX_CHART_MAX_ROWS and s.channels:
        ws_data, rows = data_sheets[0]
        x_col = 1
        for name in (t("col_time_s"), t("col_elapsed"), t("col_seq")):
            if name in heads:
                x_col = heads.index(name) + 1
                break
        first_ch = len(heads) - len(s.channels) + 1
        ws_chart = wb.create_sheet(t("sheet_chart"))
        ch = ScatterChart()
        ch.title = t("chart_title")
        ch.style = 13
        ch.height, ch.width = 12, 26
        ch.x_axis.title = heads[x_col - 1]
        ch.y_axis.title = t("value")
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        xref = Reference(ws_data, min_col=x_col, min_row=2, max_row=rows + 1)
        palette = ["1F6FEB", "2DA44E", "CF222E", "8250DF", "BF8700", "1B7C83", "E16F24", "57606A"]
        for ci in range(first_ch, len(heads) + 1):
            yref = Reference(ws_data, min_col=ci, min_row=1, max_row=rows + 1)
            ser = Series(yref, xref, title_from_data=True)
            ser.marker.symbol = "none"
            ser.smooth = False
            ser.graphicalProperties.line.width = 12700
            ser.graphicalProperties.line.solidFill = palette[(ci - first_ch) % len(palette)]
            ch.series.append(ser)
        ws_chart.add_chart(ch, "A1")

    wb.save(path)
