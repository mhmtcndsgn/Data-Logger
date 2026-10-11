"""Daha önce kaydedilmiş dosyaları (yeni TXT/CSV/XLSX ya da eski sürüm .txt) açar."""
from __future__ import annotations

import os

from .i18n import STRINGS
from .parser import parse_line, strip_old_timestamp, NUM_RE
from .session import Session, _dedupe

_META_KEYS = ("col_index", "col_seq", "col_time_s", "col_pc_time", "col_elapsed")


def _meta_names() -> dict[str, str]:
    out = {}
    for lang in STRINGS.values():
        for k in _META_KEYS:
            out[lang[k].strip().lower()] = k
    return out


def _to_float(v):
    if v is None or v == "":
        return float("nan")
    if isinstance(v, (int, float)):
        return float(v)
    v = str(v).strip()
    return float(v) if NUM_RE.match(v) else float("nan")


def _session_from_table(heads: list[str], rows, channel_prefix: str) -> Session | None:
    meta = _meta_names()
    roles = [meta.get(str(h).strip().lower()) for h in heads]
    if not any(roles):
        return None
    s = Session(mode="file", seq_mode="off", channel_prefix=channel_prefix)
    ch_idx = [i for i, r in enumerate(roles) if r is None]
    names = _dedupe([str(heads[i]).strip() or f"{channel_prefix} {j + 1}" for j, i in enumerate(ch_idx)])
    seq_i = roles.index("col_seq") if "col_seq" in roles else None
    time_i = roles.index("col_time_s") if "col_time_s" in roles else None
    el_i = roles.index("col_elapsed") if "col_elapsed" in roles else None
    for name in names:
        s._ensure(name)
    s.has_seq = seq_i is not None
    times = []
    for r in rows:
        if not r or all(c in (None, "") for c in r):
            continue
        r = list(r) + [None] * (len(heads) - len(r))
        for j, i in enumerate(ch_idx):
            s.cols[j].append(_to_float(r[i]))
        if seq_i is not None:
            v = _to_float(r[seq_i])
            sv = int(v) if v == v else s.n
            s.seq.append(s._check_seq(sv))
        tv = _to_float(r[el_i]) if el_i is not None else float(s.n)
        s.t_host.append(tv if tv == tv else float(s.n))
        if time_i is not None and len(times) < 2:
            times.append(_to_float(r[time_i]))
        s.n += 1
    if len(times) == 2 and times[1] > times[0]:
        s.sample_rate = round(1.0 / (times[1] - times[0]), 6)
    if el_i is not None:
        s.mode = "live"   # x ekseni PC zamanı
    return s


def load_file(path: str, channel_prefix: str = "Kanal") -> Session:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xlsx":
        s = _load_xlsx(path, channel_prefix)
    else:
        s = _load_text(path, channel_prefix)
    s.port = os.path.basename(path)
    s.finish("file")
    return s


def _load_xlsx(path: str, channel_prefix: str) -> Session:
    from openpyxl import load_workbook

    import itertools

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = []
        for ws in wb.worksheets:
            it = ws.iter_rows(values_only=True)
            try:
                heads = [str(h) if h is not None else "" for h in next(it)]
            except StopIteration:
                continue
            sheets.append((heads, it))
        for k, (heads, it) in enumerate(sheets):
            # >1M satırlık kayıtlar birden çok "Veriler" sayfasına bölünür: aynı başlıklıları birleştir
            rest = [it2 for h2, it2 in sheets[k + 1:] if h2 == heads]
            s = _session_from_table(heads, itertools.chain(it, *rest), channel_prefix)
            if s is not None:
                return s
    finally:
        wb.close()
    raise ValueError("no-data")


def _load_text(path: str, channel_prefix: str) -> Session:
    with open(path, "r", encoding="utf-8-sig", errors="replace") as f:
        lines = [ln.rstrip("\r\n") for ln in f]
    body = [ln for ln in lines if ln.strip() and not ln.lstrip().startswith("#")]
    if body:
        first = body[0]
        delim = "\t" if "\t" in first else ("," if "," in first else None)
        if delim:
            heads = [h.strip() for h in first.split(delim)]
            rows = ([c.strip() for c in ln.split(delim)] for ln in body[1:])
            s = _session_from_table(heads, rows, channel_prefix)
            if s is not None:
                return s
    # Genel çözümleyici (eski sürüm kayıtları, ham yedekler, düz loglar)
    s = Session(mode="file", seq_mode="auto", channel_prefix=channel_prefix)
    for i, ln in enumerate(lines):
        p = parse_line(strip_old_timestamp(ln))
        if p.kind in ("empty",):
            continue
        s.add(float(i), p)
    s.flush_pending()
    if s.n == 0:
        raise ValueError("no-data")
    return s
